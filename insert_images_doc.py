import os
import docx
from docx import Document
from docx.shared import Inches

doc_path = r'C:/Users/hp/UTS-RPL/Lembar_Jawaban_UTS_RPL_240104013_Yusuf_Rizal_Indrianto.docx'
img_dir = r'C:/Users/hp/UTS-RPL/images'

doc = Document(doc_path)

# 1. Table 4: Bukti halaman PHP berjalan
cell_t4 = doc.tables[4].rows[1].cells[0]
p4 = cell_t4.add_paragraph()
r4 = p4.add_run()
img4_path = os.path.join(img_dir, "bukti_halaman_php.png")
if os.path.exists(img4_path):
    r4.add_picture(img4_path, width=Inches(5.8))
    print("Inserted bukti_halaman_php.png into Table 4")

# 2. Table 8: Bukti riwayat commit dan pemulihan
cell_t8 = doc.tables[8].rows[1].cells[0]
p8 = cell_t8.add_paragraph()
r8 = p8.add_run()
img8_path = os.path.join(img_dir, "bukti_git_commit_restore.png")
if os.path.exists(img8_path):
    r8.add_picture(img8_path, width=Inches(5.8))
    print("Inserted bukti_git_commit_restore.png into Table 8")

# 3. Table 12: Bukti feature branch, pull request, dan konflik
cell_t12 = doc.tables[12].rows[1].cells[0]
p12 = cell_t12.add_paragraph()
r12 = p12.add_run()
img12_path = os.path.join(img_dir, "bukti_github_pr_conflict.png")
if os.path.exists(img12_path):
    r12.add_picture(img12_path, width=Inches(5.8))
    print("Inserted bukti_github_pr_conflict.png into Table 12")

# 4. Table 16: Diagram alur SDLC
cell_t16 = doc.tables[16].rows[1].cells[0]
p16 = cell_t16.add_paragraph()
r16 = p16.add_run()
img16_path = os.path.join(img_dir, "diagram_sdlc_waterfall.png")
if os.path.exists(img16_path):
    r16.add_picture(img16_path, width=Inches(5.8))
    print("Inserted diagram_sdlc_waterfall.png into Table 16")

doc.save(doc_path)
print("Updated Word document with inserted pictures successfully!")
